import json
import os
import re
import time
import threading
from collections import deque
from urllib.parse import urljoin, urlparse

import requests
from io import BytesIO
import xlrd
from openpyxl import load_workbook
from bs4 import BeautifulSoup
from flask import Flask, jsonify

BASE_URL = os.getenv("OPEN24_BASE_URL", "https://b2b.open24.lt/").rstrip("/") + "/"
USERNAME = os.getenv("OPEN24_USERNAME", "")
PASSWORD = os.getenv("OPEN24_PASSWORD", "")
QUERY = os.getenv("OPEN24_QUERY", "").strip()
MAX_PAGES = int(os.getenv("OPEN24_MAX_PAGES", "60"))
MAX_RESULTS = int(os.getenv("OPEN24_MAX_RESULTS", "150"))
TIMEOUT = int(os.getenv("OPEN24_TIMEOUT", "30"))

app = Flask(__name__)
STATE = {
    "configured": bool(USERNAME and PASSWORD),
    "status": "starting",
    "query": QUERY,
    "results": [],
    "pages_visited": 0,
    "error": None,
}

PRICE_RE = re.compile(r"(?:€\s*|EUR\s*)?\d{1,4}(?:[\.,]\d{2})?\s*(?:€|EUR)?", re.I)


def clean_text(s):
    return re.sub(r"\s+", " ", s or "").strip()


def same_host(url):
    return urlparse(url).netloc == urlparse(BASE_URL).netloc


def find_login_form(soup):
    for form in soup.find_all("form"):
        if form.find("input", {"type": "password"}):
            return form
    return None


def login(session):
    r = session.get(BASE_URL, timeout=TIMEOUT)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "lxml")
    form = find_login_form(soup)

    if not form:
        # Could already be authenticated or the login lives elsewhere.
        return r.url, soup

    payload = {}
    for inp in form.find_all("input"):
        name = inp.get("name")
        if not name:
            continue
        typ = (inp.get("type") or "text").lower()
        if typ in {"hidden", "submit"}:
            payload[name] = inp.get("value", "")

    pwd = form.find("input", {"type": "password"})
    payload[pwd.get("name")] = PASSWORD

    user_input = None
    for inp in form.find_all("input"):
        typ = (inp.get("type") or "text").lower()
        if typ in {"email", "text"} and inp.get("name"):
            user_input = inp
            break
    if not user_input:
        raise RuntimeError("Could not identify username/email field in login form.")
    payload[user_input.get("name")] = USERNAME

    action = urljoin(r.url, form.get("action") or r.url)
    method = (form.get("method") or "post").lower()
    if method == "get":
        rr = session.get(action, params=payload, timeout=TIMEOUT)
    else:
        rr = session.post(action, data=payload, timeout=TIMEOUT, allow_redirects=True)
    rr.raise_for_status()

    ss = BeautifulSoup(rr.text, "lxml")
    if find_login_form(ss):
        title = clean_text(ss.title.get_text(" ", strip=True) if ss.title else "")
        raise RuntimeError(f"Login appears to have failed; login form still present. Page title: {title}")
    return rr.url, ss


def product_candidates(soup, page_url):
    seen = set()
    out = []

    selectors = [
        "article",
        ".product",
        ".product-item",
        ".product-card",
        ".product-miniature",
        ".card",
        "tr",
    ]
    nodes = []
    for sel in selectors:
        nodes.extend(soup.select(sel))

    for node in nodes:
        text = clean_text(node.get_text(" ", strip=True))
        if len(text) < 8:
            continue
        if not PRICE_RE.search(text):
            continue

        link = node.find("a", href=True)
        href = urljoin(page_url, link["href"]) if link else page_url
        name = ""
        for sel in [".product-title", ".name", ".title", "h1", "h2", "h3", "h4", "a"]:
            el = node.select_one(sel)
            if el:
                name = clean_text(el.get_text(" ", strip=True))
                if name:
                    break
        if not name:
            name = text[:180]

        prices = [clean_text(x) for x in PRICE_RE.findall(text)]
        key = (name.casefold(), href)
        if key in seen:
            continue
        seen.add(key)

        out.append({
            "name": name,
            "url": href,
            "text": text[:1200],
            "prices": prices[:5],
        })
    return out


def link_priority(href, text):
    s = (href + " " + text).casefold()
    hot = ["product", "catalog", "shop", "keen", "footwear", "collection", "category", "prekes", "prek", "shoe", "boot"]
    return any(k in s for k in hot)




def export_keen_stock(session):
    url = urljoin(BASE_URL, "catalog/keen/?download_brand_xls=12")
    try:
        r = session.get(url, timeout=TIMEOUT)
        r.raise_for_status()
        ctype = r.headers.get("content-type","")
        print("OPEN24_XLS_META " + json.dumps({
            "url": r.url,
            "content_type": ctype,
            "bytes": len(r.content),
            "head_hex": r.content[:8].hex(),
        }, ensure_ascii=False), flush=True)

        rows = []
        data = r.content
        if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            wb = xlrd.open_workbook(file_contents=data)
            for sname in wb.sheet_names():
                sh = wb.sheet_by_name(sname)
                print("OPEN24_XLS_SHEET " + json.dumps({"sheet": sname, "rows": sh.nrows, "cols": sh.ncols}, ensure_ascii=False), flush=True)
                for i in range(sh.nrows):
                    vals = [clean_text(str(sh.cell_value(i,j))) for j in range(sh.ncols)]
                    if any(vals):
                        print("OPEN24_XLS_ROW " + json.dumps({"sheet": sname, "row": i, "cells": vals}, ensure_ascii=False), flush=True)
        else:
            wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
            for ws in wb.worksheets:
                print("OPEN24_XLS_SHEET " + json.dumps({"sheet": ws.title, "rows": ws.max_row, "cols": ws.max_column}, ensure_ascii=False), flush=True)
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    vals = [clean_text("" if v is None else str(v)) for v in row]
                    if any(vals):
                        print("OPEN24_XLS_ROW " + json.dumps({"sheet": ws.title, "row": i, "cells": vals}, ensure_ascii=False), flush=True)
    except Exception as e:
        print("OPEN24_XLS_ERROR " + json.dumps({"error": str(e)}, ensure_ascii=False), flush=True)


def diagnose_order(session, landing_url):
    keywords = ["order", "orders", "cart", "basket", "checkout", "užsak", "uzsak", "krep", "pirk", "mano"]
    seen = set()
    queue = [landing_url, BASE_URL]
    candidates = []
    for seed in queue:
        try:
            r = session.get(seed, timeout=TIMEOUT)
            r.raise_for_status()
            soup = BeautifulSoup(r.text, "lxml")
            print("OPEN24_ORDER_PAGE " + json.dumps({
                "url": r.url,
                "title": clean_text(soup.title.get_text(" ", strip=True) if soup.title else ""),
                "text": clean_text(soup.get_text(" ", strip=True))[:2500],
            }, ensure_ascii=False), flush=True)
            for a in soup.find_all("a", href=True):
                txt = clean_text(a.get_text(" ", strip=True))
                href = urljoin(r.url, a.get("href"))
                hay = (txt + " " + href).casefold()
                if same_host(href) and any(k in hay for k in keywords):
                    candidates.append((txt, href))
        except Exception as e:
            print("OPEN24_ORDER_SEED_ERROR " + json.dumps({"url": seed, "error": str(e)}, ensure_ascii=False), flush=True)

    for txt, href in candidates:
        if href in seen:
            continue
        seen.add(href)
        try:
            r = session.get(href, timeout=TIMEOUT)
            r.raise_for_status()
            soup = BeautifulSoup(r.text, "lxml")
            print("OPEN24_ORDER_CANDIDATE " + json.dumps({
                "link_text": txt,
                "url": r.url,
                "title": clean_text(soup.title.get_text(" ", strip=True) if soup.title else ""),
                "text": clean_text(soup.get_text(" ", strip=True))[:5000],
            }, ensure_ascii=False), flush=True)
            for ti, table in enumerate(soup.find_all("table")):
                for ri, tr in enumerate(table.find_all("tr")):
                    cells = [clean_text(x.get_text(" ", strip=True)) for x in tr.find_all(["th","td"])]
                    if any(cells):
                        print("OPEN24_ORDER_ROW " + json.dumps({
                            "url": r.url, "table": ti, "row": ri, "cells": cells
                        }, ensure_ascii=False), flush=True)
            for fi, form in enumerate(soup.find_all("form")):
                fields=[]
                for inp in form.find_all(["input","select","textarea","button"]):
                    fields.append({
                        "tag": inp.name,
                        "name": inp.get("name"),
                        "type": inp.get("type"),
                        "value": inp.get("value"),
                        "text": clean_text(inp.get_text(" ", strip=True))[:200],
                    })
                print("OPEN24_ORDER_FORM " + json.dumps({
                    "url": r.url,
                    "index": fi,
                    "action": urljoin(r.url, form.get("action") or r.url),
                    "method": (form.get("method") or "get").lower(),
                    "fields": fields[:80],
                }, ensure_ascii=False), flush=True)
        except Exception as e:
            print("OPEN24_ORDER_CANDIDATE_ERROR " + json.dumps({"url": href, "error": str(e)}, ensure_ascii=False), flush=True)


def diagnose_keen(session):
    url = urljoin(BASE_URL, "catalog/keen/")
    try:
        r = session.get(url, timeout=TIMEOUT)
        r.raise_for_status()
        soup = BeautifulSoup(r.text, "lxml")
        print("OPEN24_DIAG_PAGE " + json.dumps({
            "url": r.url,
            "title": clean_text(soup.title.get_text(" ", strip=True) if soup.title else ""),
        }, ensure_ascii=False), flush=True)

        for i, a in enumerate(soup.find_all("a", href=True)[:300]):
            txt = clean_text(a.get_text(" ", strip=True))
            href = urljoin(r.url, a.get("href"))
            if txt:
                print("OPEN24_DIAG_LINK " + json.dumps({"i": i, "text": txt[:300], "href": href}, ensure_ascii=False), flush=True)

        for ti, table in enumerate(soup.find_all("table")):
            for ri, tr in enumerate(table.find_all("tr")):
                cells = [clean_text(x.get_text(" ", strip=True)) for x in tr.find_all(["th","td"])]
                if any(cells):
                    print("OPEN24_DIAG_ROW " + json.dumps({"table": ti, "row": ri, "cells": cells}, ensure_ascii=False), flush=True)
    except Exception as e:
        print("OPEN24_DIAG_ERROR " + json.dumps({"error": str(e)}, ensure_ascii=False), flush=True)


def crawl(session, start_url):
    q = deque([start_url])
    visited = set()
    products = {}

    while q and len(visited) < MAX_PAGES:
        url = q.popleft()
        if url in visited or not same_host(url):
            continue
        visited.add(url)

        try:
            r = session.get(url, timeout=TIMEOUT)
            if r.status_code >= 400:
                continue
            if "text/html" not in r.headers.get("content-type", ""):
                continue
        except requests.RequestException:
            continue

        soup = BeautifulSoup(r.text, "lxml")
        for p in product_candidates(soup, r.url):
            key = (p["name"].casefold(), p["url"])
            products[key] = p

        links = []
        for a in soup.find_all("a", href=True):
            href = urljoin(r.url, a["href"])
            if not same_host(href):
                continue
            if href.startswith("mailto:") or href.startswith("javascript:"):
                continue
            text = clean_text(a.get_text(" ", strip=True))
            if link_priority(href, text):
                links.append(href)

        for href in links[:80]:
            if href not in visited:
                q.append(href)

        time.sleep(0.15)

    return list(products.values()), len(visited)


def matches_query(item, query):
    if not query:
        return True
    hay = " ".join([item.get("name", ""), item.get("text", ""), item.get("url", "")]).casefold()
    terms = [t for t in re.split(r"\s+", query.casefold()) if t]
    return all(t in hay for t in terms)


def run_sync():
    if not USERNAME or not PASSWORD:
        STATE.update(status="not_configured", error="OPEN24_USERNAME / OPEN24_PASSWORD are not set")
        print("OPEN24_STATUS " + json.dumps(STATE, ensure_ascii=False), flush=True)
        return

    session = requests.Session()
    session.headers.update({
        "User-Agent": "Mozilla/5.0 (compatible; Outfish-B2B/1.0; +https://outfish.lv/)",
        "Accept-Language": "en-US,en;q=0.9,lt;q=0.8",
    })

    try:
        landing, _ = login(session)
        if QUERY.casefold() == "keen":
            diagnose_keen(session)
            export_keen_stock(session)
        if QUERY.casefold() == "__order__":
            diagnose_order(session, landing)
        products, pages = crawl(session, landing)
        filtered = [p for p in products if matches_query(p, QUERY)]
        filtered = filtered[:MAX_RESULTS]
        STATE.update(status="ok", results=filtered, pages_visited=pages, error=None)
        print("OPEN24_STATUS " + json.dumps({
            "status": "ok",
            "query": QUERY,
            "pages_visited": pages,
            "products_found": len(products),
            "results_returned": len(filtered),
        }, ensure_ascii=False), flush=True)
        for item in filtered:
            print("OPEN24_RESULT " + json.dumps(item, ensure_ascii=False), flush=True)
    except Exception as e:
        STATE.update(status="error", error=str(e))
        print("OPEN24_ERROR " + json.dumps({"error": str(e)}, ensure_ascii=False), flush=True)


def _background_start():
    STATE.update(status="running", error=None)
    run_sync()

threading.Thread(target=_background_start, daemon=True).start()


@app.get("/health")
def health():
    return jsonify({
        "status": STATE["status"],
        "configured": STATE["configured"],
        "query": STATE["query"],
        "pages_visited": STATE["pages_visited"],
        "results": len(STATE["results"]),
        "error": STATE["error"],
    })


@app.get("/")
def root():
    return jsonify({"service": "outfish-open24-b2b", "status": STATE["status"], "query": STATE["query"], "results": len(STATE["results"])})


if __name__ == "__main__":
    port = int(os.getenv("PORT", "10000"))
    app.run(host="0.0.0.0", port=port)
